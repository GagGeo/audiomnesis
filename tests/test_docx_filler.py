import docx
import pytest

from anamnese import config
from anamnese.docx_filler import fill_report
from anamnese.sections import norm

SECTION_TEXTS = {
    "plaintes": "Spontanément, Monsieur XXX mentionne une plainte A.\nSur demande, il précise B.",
    "atcd_perso": "Interrogé sur ses antécédents, Monsieur XXX rapporte C.",
    "atcd_fam": "Concernant ses antécédents familiaux, Monsieur XXX mentionne D.",
    "developpement": "Concernant son développement, E.\nDurant sa scolarité, F.\nSur le plan professionnel, G.",
    "thymie": None,  # section that failed / had no info: must stay as template stubs
}
HEADER = {"laterality": "Droitier", "langue_maternelle": "Français",
          "profession": "Enseignant", "accompagnant": None}


@pytest.fixture
def filled(tmp_path):
    out = tmp_path / "out.docx"
    fill_report(config.DEFAULT_TEMPLATE, out,
                intro_text="Je rencontre ce jour Monsieur XXX pour un bilan.",
                section_texts=SECTION_TEXTS, header_fields=HEADER)
    return docx.Document(str(out))


def _body(doc):
    return [p.text for p in doc.paragraphs]


def test_sections_filled_with_multiple_paragraphs(filled):
    body = _body(filled)
    assert "Spontanément, Monsieur XXX mentionne une plainte A." in body
    assert "Sur demande, il précise B." in body
    assert "Durant sa scolarité, F." in body


def test_failed_section_keeps_template_stubs(filled):
    assert any(norm(t).startswith("d'un point de vue thymique") for t in _body(filled))


def test_intro_filled(filled):
    assert "Je rencontre ce jour Monsieur XXX pour un bilan." in _body(filled)


def test_conclusions_untouched(filled):
    body = _body(filled)
    assert any(norm(t) == norm("L’examen de met en évidence ") for t in body)
    assert any(norm(t) == norm("Dans ce contexte, nous recommandons") for t in body)


def test_header_fields_filled_only_when_present(filled):
    cells = "\n".join(c.text for r in filled.tables[1].rows for c in r.cells)
    cells = cells.replace("\xa0", " ")
    assert "Latéralité : Droitier" in cells
    assert "Profession : Enseignant" in cells.replace("  ", " ")
    assert "None" not in cells  # absent fields left blank, never the string "None"


def test_original_template_is_never_modified(tmp_path):
    import hashlib
    before = hashlib.sha256(config.DEFAULT_TEMPLATE.read_bytes()).hexdigest()
    fill_report(config.DEFAULT_TEMPLATE, tmp_path / "x.docx",
                intro_text="Test.", section_texts=SECTION_TEXTS,
                header_fields=HEADER)
    after = hashlib.sha256(config.DEFAULT_TEMPLATE.read_bytes()).hexdigest()
    assert before == after
