"""Fill a copy of the report template with generated text, preserving layout."""
import copy
from pathlib import Path

import docx

from anamnese import sections
from anamnese.sections import SECTIONS, norm

# label prefix (normalized) -> header_fields key, for document.tables[1],
# per report language (matches the FR / EN template's header table).
_HEADER_LABELS = {
    "fr": {
        "latéralité :": "laterality",
        "langue maternelle :": "langue_maternelle",
        "statut marital :": "statut_marital",
        "niveau d'études :": "niveau_etudes",
        "profession :": "profession",
        "motif de consultation :": "motif_consultation",
        "accompagnant :": "accompagnant",
    },
    "en": {
        "handedness:": "laterality",
        "mother tongue:": "langue_maternelle",
        "marital status:": "statut_marital",
        "level of education:": "niveau_etudes",
        "profession:": "profession",
        "reason for consultation:": "motif_consultation",
        "support person:": "accompagnant",
    },
}


def _set_text(par, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    if par.runs:
        par.runs[0].text = text
        for run in par.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        par.add_run(text)


def _insert_paragraph_after(par, text):
    """Clone `par` (same style/formatting), set its text, insert after `par`."""
    new_el = copy.deepcopy(par._p)
    par._p.addnext(new_el)
    new_par = docx.text.paragraph.Paragraph(new_el, par._parent)
    _set_text(new_par, text)
    return new_par


def _find_paragraph(doc, predicate):
    for i, p in enumerate(doc.paragraphs):
        if predicate(p.text):
            return i, p
    return None, None


def _section_stub_paragraphs(doc, section, lang):
    """Stub paragraphs = paragraphs after the heading until the next heading."""
    headings = {norm(sections.heading(s, lang)) for s in SECTIONS}
    target = sections.heading(section, lang)
    idx, _ = _find_paragraph(doc, lambda t: norm(t) == norm(target))
    if idx is None:
        raise ValueError(f"Heading not found: {target}")
    stubs = []
    for p in doc.paragraphs[idx + 1:]:
        if norm(p.text) in headings:
            break
        if p.text.strip():
            stubs.append(p)
    return stubs


def _fill_section(doc, section, text, lang):
    stubs = _section_stub_paragraphs(doc, section, lang)
    if not stubs:
        raise ValueError(f"No stub paragraphs under: {sections.heading(section, lang)}")
    paras = [t.strip() for t in text.split("\n") if t.strip()]
    # remove extra stubs beyond the first, then grow from the first
    anchor = stubs[0]
    for p in stubs[1:]:
        p._p.getparent().remove(p._p)
    _set_text(anchor, paras[0])
    for t in paras[1:]:
        anchor = _insert_paragraph_after(anchor, t)


def _fill_header_fields(doc, header_fields, lang):
    table = doc.tables[1]
    labels = _HEADER_LABELS[lang]
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            for p in cell.paragraphs:
                n = norm(p.text)
                for label, key in labels.items():
                    value = header_fields.get(key)
                    if value and n == label.rstrip():
                        _set_text(p, p.text.rstrip() + " " + str(value))


def fill_report(template_path: Path, output_path: Path, intro_text: str,
                section_texts: dict, header_fields: dict, lang: str = "fr") -> None:
    """Write output_path from template_path. Never modifies the template.

    section_texts: section key -> generated text (paragraphs separated by
    newline) or None to leave the template stubs untouched.
    lang: report language, selecting which heading/label anchors to match.
    """
    doc = docx.Document(str(template_path))
    if intro_text:
        leadin = norm(sections.leadin(lang))
        _, intro_par = _find_paragraph(doc, lambda t: norm(t).startswith(leadin))
        if intro_par is not None:
            _set_text(intro_par, intro_text)
    for section in SECTIONS:
        text = section_texts.get(section.key)
        if text:
            _fill_section(doc, section, text, lang)
    if header_fields:
        _fill_header_fields(doc, header_fields, lang)
    doc.save(str(output_path))
